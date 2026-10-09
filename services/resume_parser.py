"""
Resume Parsing Service for extracting raw text from PDF, DOCX, DOC, and TXT files.
"""
import io
import os
import re
from typing import Dict, Any

from utils.validators import sanitize_text

# Import extraction backends safely
try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    import pypdf
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF bytes using pdfplumber with pypdf fallback."""
    extracted_text = ""
    
    # Try pdfplumber first
    if HAS_PDFPLUMBER:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                page_texts = []
                for page in pdf.pages:
                    txt = page.extract_text()
                    if txt:
                        page_texts.append(txt)
                extracted_text = "\n\n".join(page_texts)
        except Exception:
            extracted_text = ""
            
    # Fallback to pypdf if pdfplumber extracted nothing or isn't present
    if not extracted_text.strip() and HAS_PYPDF:
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            page_texts = []
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    page_texts.append(txt)
            extracted_text = "\n\n".join(page_texts)
        except Exception as e:
            if not extracted_text:
                raise RuntimeError(f"PDF extraction failed: {str(e)}")

    return extracted_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX bytes using python-docx."""
    if not HAS_DOCX:
        raise ImportError("python-docx package is not installed.")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        
        # Paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
                
        # Tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
                    
        return "\n".join(full_text)
    except Exception as e:
        raise RuntimeError(f"DOCX extraction error: {str(e)}")


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extracts text from raw TXT bytes using multi-encoding fallback."""
    for encoding in ["utf-8", "latin-1", "cp1252"]:
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def parse_resume_file(uploaded_file) -> Dict[str, Any]:
    """
    Main entry point for parsing uploaded resume files.
    Returns structured dictionary with extraction results.
    """
    if uploaded_file is None:
        return {
            "text": "",
            "file_name": "",
            "file_size": 0,
            "file_type": "",
            "success": False,
            "error_message": "No file provided.",
            "word_count": 0
        }

    file_name = uploaded_file.name
    file_bytes = uploaded_file.getvalue()
    file_size = len(file_bytes)
    _, ext = os.path.splitext(file_name.lower())

    extracted_text = ""
    error_message = None

    try:
        if ext == ".pdf":
            extracted_text = extract_text_from_pdf(file_bytes)
        elif ext == ".docx":
            extracted_text = extract_text_from_docx(file_bytes)
        elif ext == ".doc":
            # Attempt docx parsing or fallback text decoding for .doc
            try:
                extracted_text = extract_text_from_docx(file_bytes)
            except Exception:
                extracted_text = extract_text_from_txt(file_bytes)
                # Filter printable ascii/utf
                extracted_text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', extracted_text)
        elif ext == ".txt":
            extracted_text = extract_text_from_txt(file_bytes)
        else:
            error_message = f"Unsupported file extension: {ext}"
    except Exception as e:
        error_message = f"Failed to extract text from file: {str(e)}"

    sanitized = sanitize_text(extracted_text)
    word_count = len(sanitized.split()) if sanitized else 0

    if not sanitized and not error_message:
        error_message = "No readable text could be extracted from the uploaded document. It might be scanned/image-based."

    success = bool(sanitized) and not error_message

    return {
        "text": sanitized,
        "file_name": file_name,
        "file_size": file_size,
        "file_type": ext.replace(".", "").upper(),
        "success": success,
        "error_message": error_message,
        "word_count": word_count
    }
